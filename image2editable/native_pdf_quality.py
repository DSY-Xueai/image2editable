"""Validate native PDF output against its source objects and rendered pixels."""

from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import shutil

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.oxml.ns import qn

from image2editable.component_quality import evaluate_page_quality
from image2editable.inputs import sha256_file
from scripts.visual_segment import visual_difference


def evaluate_native_pdf_page(
    *, pptx_path: Path, page_number: int, analysis: dict,
    source_path: Path, rendered_path: Path,
) -> dict:
    presentation = Presentation(pptx_path)
    slide = presentation.slides[page_number - 1]
    objects = sorted(analysis["objects"], key=lambda item: item["z_index"])
    names = [f"image2editable:{item['id']}" for item in objects]
    actual_names = [shape.name for shape in slide.shapes]
    structure_ok = actual_names == names and len(set(names)) == len(names)
    text_ok = structure_ok
    object_reports = []
    by_name = {shape.name: shape for shape in slide.shapes}
    source_width = analysis["width_pt"]
    source_height = analysis["height_pt"]
    scale = min(presentation.slide_width / source_width, presentation.slide_height / source_height)
    offset_x = (presentation.slide_width - source_width * scale) / 2
    offset_y = (presentation.slide_height - source_height * scale) / 2
    for item, name in zip(objects, names, strict=True):
        shape = by_name.get(name)
        violations = []
        if shape is None:
            violations.append("missing_object")
        else:
            left, bottom, right, top = item["bbox_pt"]
            expected = [
                offset_x + left * scale, offset_y + (source_height - top) * scale,
                (right - left) * scale, (top - bottom) * scale,
            ]
            actual = [shape.left, shape.top, shape.width, shape.height]
            if any(abs(a - b) > 2 for a, b in zip(actual, expected, strict=True)):
                violations.append("object_geometry")
            if item["type"] == "text":
                if not shape.has_text_frame or shape.text != item["text"]:
                    violations.append("editable_text")
                    text_ok = False
                else:
                    runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
                    if not runs or any(
                        run.font.size is None
                        or abs(run.font.size - item["font_size"] * scale) > 127
                        or bool(run.font.bold) != bool(item.get("bold"))
                        or bool(run.font.italic) != bool(item.get("italic"))
                        or run.font.color.type != MSO_COLOR_TYPE.RGB
                        or tuple(run.font.color.rgb) != tuple(item["color_rgb"])
                        or run.font.name != str(item.get("font") or "Arial").lstrip("/")
                        for run in runs
                    ):
                        violations.append("editable_text_style")
                    for run in runs:
                        alpha = run._r.find(".//" + qn("a:alpha"))
                        opacity = 100000 if alpha is None else int(alpha.get("val"))
                        if abs(opacity - round(item.get("fill_opacity", 1) * 100000)) > 1:
                            violations.append("editable_text_opacity")
            elif item["type"] in {"image", "patch"}:
                if shape.shape_type != 13:
                    violations.append("image_object_type")
                else:
                    if hashlib.sha256(shape.image.blob).hexdigest() != item["asset_sha256"]:
                        violations.append("image_asset_mismatch")
            elif shape.shape_type not in {1, 9}:
                violations.append("native_shape_type")
        object_reports.append({
            "component_id": item["id"], "accepted": not violations,
            "violations": violations, "metrics": {},
        })
    expected_text = Counter(item["text"] for item in objects if item["type"] == "text")
    actual_text = Counter(shape.text for shape in slide.shapes if shape.has_text_frame and shape.text)
    text_ok = text_ok and expected_text == actual_text
    with Image.open(source_path) as image:
        source_image = image.convert("RGB")
    with Image.open(rendered_path) as image:
        rendered = np.asarray(image.convert("RGB"))
    source = _source_canvas(source_image, (rendered.shape[1], rendered.shape[0]))
    metrics = visual_difference(source, rendered, np.zeros(source.shape[:2], dtype=np.uint8))
    report = evaluate_page_quality(
        object_reports, visual_metrics=metrics,
        page_checks={"pptx_reopen": "pass", "editable_text_once": "pass" if text_ok else "fail"},
        expected_component_ids=[item["id"] for item in objects],
        initial_component_count=len(objects), active_visual_count=len(objects),
    )
    if not structure_ok:
        report["violations"].append("native_object_order")
        report["accepted"] = False
    return {
        "schema_version": 1, "page_number": page_number,
        "pptx_sha256": sha256_file(pptx_path),
        "source_sha256": sha256_file(source_path),
        "rendered_sha256": sha256_file(rendered_path),
        "report": report,
    }


def _source_canvas(source: Image.Image, size: tuple[int, int]) -> np.ndarray:
    if source.size == size:
        return np.asarray(source)
    scale = min(size[0] / source.width, size[1] / source.height)
    dimensions = (round(source.width * scale), round(source.height * scale))
    with Image.new("RGB", size, "white") as canvas:
        with source.resize(dimensions, Image.Resampling.LANCZOS) as content:
            canvas.paste(content, ((size[0] - dimensions[0]) // 2, (size[1] - dimensions[1]) // 2))
        return np.asarray(canvas).copy()


def _native_renderers():
    from image2editable.libreoffice_renderer import LibreOfficeRenderer
    from image2editable.powerpoint_renderer import PowerPointRenderer

    return (PowerPointRenderer.discover(), LibreOfficeRenderer.discover())


def validate_native_pdf_output(store, page_records, pptx_path: Path, variant: str) -> dict:
    """Validate the staged, final deck before any output is published."""
    from image2editable.legacy import _load_legacy_ref
    from image2editable.powerpoint_renderer import RendererUnavailable

    native_pages = [
        (number, page_id, state)
        for number, (page_id, state, _, _) in enumerate(page_records, start=1)
        if state.get("route") == "pdf_native"
    ]
    if not native_pages:
        return {}
    renderers = [renderer for renderer in _native_renderers() if renderer.available()]
    if not renderers:
        raise RendererUnavailable("Native PDF validation requires PowerPoint or LibreOffice")
    presentation = Presentation(pptx_path)
    ratio = presentation.slide_width / presentation.slide_height
    reports = {}
    for number, page_id, state in native_pages:
        _, payload = _load_legacy_ref(store, state["native_page_ref"])
        native = json.loads(payload.decode("utf-8"))
        reconstruction = store.root / "pages" / page_id / "reconstruction"
        source = reconstruction / "native-source.png"
        shutil.copyfile(reconstruction.parent / "source.png", source)
        with Image.open(source) as image:
            width = max(image.width, math.ceil(image.height * ratio))
            height = round(width / ratio)
        rendered = reconstruction / ("native-render-" + variant.replace(":", "x") + ".png")
        last_error = None
        for renderer in renderers:
            try:
                identity = renderer.render_page(pptx_path, number, rendered, width=width, height=height)
                break
            except RendererUnavailable as error:
                last_error = error
        else:
            raise RendererUnavailable("No native PDF renderer could be started") from last_error
        quality = evaluate_native_pdf_page(
            pptx_path=pptx_path, page_number=number, analysis=native["analysis"],
            source_path=source, rendered_path=rendered,
        )
        quality.update(
            page_id=page_id, analysis_sha256=state["native_page_ref"]["sha256"],
            renderer=identity["renderer"],
        )
        if not quality["report"]["accepted"]:
            store.write_json(reconstruction.relative_to(store.root) / "native-quality-failed.json", quality)
            raise RuntimeError("Native PDF output needs repair: " + ", ".join(quality["report"]["violations"]))
        reports[page_id] = quality
    return reports
