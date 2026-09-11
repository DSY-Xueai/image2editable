from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn


@pytest.fixture(params=["scripts", "skills/image-to-ppt/scripts"])
def assembler(request):
    path = Path(__file__).resolve().parents[1] / request.param / "ppt_assemble.py"
    spec = importlib.util.spec_from_file_location("styled_assembler", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_positioned_runs_are_native_editable_text_with_fill_outline_and_rotation(assembler, tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    item = {
        "box": [100, 100, 200, 100], "text": "AB", "font_size": 30,
        "runs": [
            {"text": "A", "box": [0.0, 0.2, 0.4, 0.6], "color": "#73c5ce",
             "rotation": -12.5, "font_size": 30, "outline_color": "#362315", "outline_width": 2},
            {"text": "B", "box": [0.5, 0.0, 0.4, 0.6], "color": "#ed9cab",
             "rotation": 6.5, "font_size": 32, "outline_color": "#462315", "outline_width": 3},
        ],
    }
    shape = assembler._add_textbox(
        slide, item, 960, 540, assembler.compute_slide_transform(960, 540, "original"),
    )
    assert len(shape.shapes) == 2
    assert [child.text for child in shape.shapes] == ["A", "B"]
    for child, expected in zip(shape.shapes, item["runs"]):
        run = child.text_frame.paragraphs[0].runs[0]
        assert str(run.font.color.rgb).lower() == expected["color"][1:]
        assert child.rotation == pytest.approx(expected["rotation"] % 360)
        outline = run._r.get_or_add_rPr().find(qn("a:ln"))
        assert outline.get("w") == str(expected["outline_width"] * 12700)
        assert outline.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == expected["outline_color"][1:]
    path = tmp_path / "native.pptx"
    deck.save(path)
    reopened = Presentation(path)
    assert [child.text for child in reopened.slides[0].shapes[0].shapes] == ["A", "B"]
    reopened.slides[0].shapes[0].shapes[0].text_frame.paragraphs[0].runs[0].text = "C"
    reopened.save(path)
    assert Presentation(path).slides[0].shapes[0].shapes[0].text == "C"


@pytest.mark.parametrize("runs", [[], [{"text": "Wrong", "box": [0, 0, 1, 1]}],
    [{"text": "AB", "box": [0, 0, float("nan"), 1]}]])
def test_invalid_positioned_runs_do_not_silently_drop_text(assembler, runs):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    with pytest.raises(ValueError, match="text runs"):
        assembler._add_textbox(
            slide, {"box": [0, 0, 100, 40], "text": "AB", "runs": runs},
            960, 540, assembler.compute_slide_transform(960, 540, "original"),
        )
    assert len(slide.shapes) == 0


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_positioned_text_rotates_about_line_center_and_scales_outline(assembler, rotation):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = assembler._add_textbox(
        slide, {"box": [100, 100, 100, 50], "text": "A", "rotation": rotation,
            "runs": [{"text": "A", "box": [0.25, 0.25, 0.5, 0.5], "font_size": 10,
                "outline_width": 2, "outline_color": "#112233"}]},
        960, 540, assembler.compute_slide_transform(960, 540, "original"),
        1920, 1080, 50, 100,
    )
    child = shape.shapes[0]
    assert child.rotation == rotation
    assert (child.left + child.width / 2) / 12700 == pytest.approx(100, abs=0.01)
    assert (child.top + child.height / 2) / 12700 == pytest.approx(112.5, abs=0.01)
    line = child.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().find(qn("a:ln"))
    assert line.get("w") == "12700"


def test_child_rotation_is_validated_before_creating_partial_group(assembler):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    with pytest.raises(ValueError, match="text runs"):
        assembler._add_textbox(slide, {
            "box": [0, 0, 100, 100], "text": "AB", "runs": [
                {"text": "A", "box": [0, 0, 0.5, 1]},
                {"text": "B", "box": [0.5, 0, 0.5, 1], "rotation": float("nan")},
            ]}, 960, 540, assembler.compute_slide_transform(960, 540, "original"))
    assert len(slide.shapes) == 0


@pytest.mark.parametrize("text", ["\uff01", "\u4e00", "A"])
@pytest.mark.parametrize("rotation", [0, 90])
@pytest.mark.parametrize("canvas_width", [None, 1920])
def test_ink_runs_preserve_font_size_and_rotate_about_visible_center(assembler, text, rotation, canvas_width):
    font = assembler._ocr_measurement_font("Microsoft YaHei", True, False)
    if font is None:
        pytest.skip("Requires Microsoft YaHei")
    mask, offset = font.getmask2(text)
    x0, y0, x1, y1 = mask.getbbox()
    ink = [x0 + offset[0], y0 + offset[1], x1 + offset[0], y1 + offset[1]]
    size, stroke = 60, 2
    width = (ink[2] - ink[0]) * size / 1000 + stroke
    height = (ink[3] - ink[1]) * size / 1000 + stroke
    deck = Presentation()
    shape = assembler._add_textbox(
        deck.slides.add_slide(deck.slide_layouts[6]),
        {"box": [100, 100, 200, 150], "text": text, "runs": [{
            "text": text, "box": [0, 0, width / 200, height / 150],
            "box_kind": "ink", "font": "Microsoft YaHei", "bold": True,
            "font_size": size, "rotation": rotation,
            "outline_width": stroke, "outline_color": "#112233",
        }]}, 960, 540, assembler.compute_slide_transform(960, 540, "original"),
        canvas_width, 1080 if canvas_width else None,
    ).shapes[0]
    scale = .5 if canvas_width else 1
    size *= scale
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == pytest.approx(size, abs=.011)
    assert shape.width.pt >= font.getlength(text) * size / 1000 - .001
    ascent, descent = font.getmetrics()
    dx = (ink[0] + ink[2]) * size / 2000 - shape.width.pt / 2
    dy = (ink[1] + ink[3] - ascent - descent) * size / 2000
    visible_x = shape.left.pt + shape.width.pt / 2 + (dx if rotation == 0 else -dy)
    visible_y = shape.top.pt + shape.height.pt / 2 + (dy if rotation == 0 else dx)
    assert visible_x == pytest.approx((100 + width / 2) * scale, abs=.02)
    assert visible_y == pytest.approx((100 + height / 2) * scale, abs=.02)


def test_invalid_box_kind_fails_before_creating_shapes(assembler):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    with pytest.raises(ValueError, match="text runs"):
        assembler._add_textbox(slide, {
            "box": [0, 0, 100, 100], "text": "A", "runs": [
                {"text": "A", "box": [0, 0, 1, 1], "box_kind": "unknown"},
            ]}, 960, 540, assembler.compute_slide_transform(960, 540, "original"))
    assert len(slide.shapes) == 0
