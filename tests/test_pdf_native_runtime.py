from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

import image_to_ppt
import image2editable.legacy as legacy
import image2editable.runtime as runtime
from image2editable.pdf_input import prepare_pdf_job
from image2editable.runtime import run_job
from image2editable.store import RunStore


@pytest.fixture(autouse=True)
def native_unit_renderer(monkeypatch):
    """Keep assembly/ownership tests independent of an installed office suite."""
    from image2editable import native_pdf_quality

    class Renderer:
        def available(self):
            return True

        def render_page(self, pptx_path, page_number, output_path, *, width, height):
            with Image.open(output_path.parent / "native-source.png") as source:
                pixels = native_pdf_quality._source_canvas(source.convert("RGB"), (width, height))
            Image.fromarray(pixels).save(output_path)
            return {"renderer": "unit-test"}

    monkeypatch.setattr(native_pdf_quality, "_native_renderers", lambda: (Renderer(),))


def _write_cropped_spacing_pdf(source: Path, image: Path) -> None:
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.saveState()
    path = document.beginPath()
    path.rect(40, 20, 25, 40)
    document.clipPath(path, stroke=0, fill=0)
    document.drawImage(str(image), 40, 20, width=30, height=40)
    document.restoreState()
    document._code.append(
        "BT /F1 12 Tf 2 Tc 1 0 0 1 80 70 Tm (Wide) Tj ET"
    )
    document.save()


def test_native_pdf_bypasses_visual_workers_and_builds_editable_pptx(
    tmp_path: Path, monkeypatch,
) -> None:
    image = tmp_path / "asset.png"
    Image.new("RGB", (12, 8), "green").save(image)
    source = tmp_path / "native.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawImage(str(image), 20, 15, width=48, height=32)
    document.setFillColorRGB(1, 0, 0)
    document.rect(75, 15, 30, 20, fill=1, stroke=0)
    document.setStrokeColorRGB(0, 0, 1)
    document.line(115, 15, 155, 15)
    document.setFillColorRGB(0, 0, 0)
    document.setFont("Helvetica", 14)
    document.drawString(80, 70, "Editable")
    document.save()

    def unexpected_worker(*args, **kwargs):
        raise AssertionError("native PDF must not create OCR or visual workers")

    monkeypatch.setattr(image_to_ppt, "create_ocr_worker_pool", unexpected_worker)
    monkeypatch.setattr(image_to_ppt, "create_visual_worker_pool", unexpected_worker)

    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
    )
    summary = run_job(run)

    output = Path(summary["outputs"]["original"])
    presentation = Presentation(output)
    slide = presentation.slides[0]
    assert len(presentation.slides) == 1
    assert [
        shape.text for shape in slide.shapes
        if shape.has_text_frame and shape.text
    ] == [
        "Editable"
    ]
    assert sum(shape.shape_type == 13 for shape in slide.shapes) == 1
    assert sum(shape.shape_type == 1 for shape in slide.shapes) == 1
    assert sum(shape.shape_type == 9 for shape in slide.shapes) == 1
    assert max(
        max(shape.image.size)
        for shape in slide.shapes
        if shape.shape_type == 13
    ) == 12
    named = [shape for shape in slide.shapes if shape.name.startswith("image2editable:")]
    assert [shape.name for shape in named] == [
        "image2editable:image-0001",
        "image2editable:shape-0001",
        "image2editable:shape-0002",
        "image2editable:text-0004",
    ]
    picture = named[0]
    from pptx.enum.text import MSO_AUTO_SIZE

    assert named[-1].text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert named[-1].text_frame.word_wrap is False
    assert abs(picture.left - Inches(1.25)) < 2
    assert abs(picture.top - Inches(4.5625)) < 2
    state = RunStore.open(run).read_json(
        "pages/page_001/reconstruction/component_state.json"
    )
    assert state["route"] == "pdf_native"
    assert state["status"] == "ready_for_assembly"
    native_path, _ = legacy._load_legacy_ref(RunStore.open(run), state["native_page_ref"])
    assert native_path.is_file()
    delivery = RunStore.open(run).read_json("pages/page_001/reconstruction/component_delivery.json")
    quality_path, _ = legacy._load_legacy_ref(RunStore.open(run), delivery["native_quality_ref"])
    assert quality_path.is_file()
    assert (quality_path.parent / "native-source.png").is_file()
    assert (quality_path.parent / "native-render-original.png").is_file()
    assert not (run / "pages/page_001/source.png").exists()
    assert not (run / "pages/page_001/pdf-assets").exists()


def test_native_pdf_rejects_replaced_extracted_asset(
    tmp_path: Path,
) -> None:
    image = tmp_path / "asset.png"
    Image.new("RGB", (12, 8), "green").save(image)
    source = tmp_path / "native.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawImage(str(image), 20, 15, width=48, height=32)
    document.drawString(80, 70, "Editable")
    document.save()
    output = tmp_path / "output.pptx"
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=output,
        slide_size="original",
    )
    request = RunStore.open(run).read_json(
        "pages/page_001/page_request.json"
    )
    picture = next(
        item for item in request["pdf_analysis"]["objects"]
        if item["type"] == "image"
    )
    (run / picture["asset_path"]).write_bytes(b"replaced")

    with pytest.raises(ValueError, match="sha256 mismatch"):
        run_job(run)

    assert not output.exists()
    assert RunStore.open(run).read_json("run_state.json")["status"] == "failed"


def test_native_pdf_assembles_from_verified_asset_bytes(
    tmp_path: Path, monkeypatch,
) -> None:
    image = tmp_path / "asset.png"
    Image.new("RGB", (12, 8), "green").save(image)
    source = tmp_path / "native.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawImage(str(image), 20, 15, width=48, height=32)
    document.save()
    output = tmp_path / "output.pptx"
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=output,
        slide_size="original",
    )
    original_load = legacy._load_legacy_ref

    def replace_after_verification(store, reference, **kwargs):
        path, payload = original_load(store, reference, **kwargs)
        if "pdf-assets" in reference["path"]:
            Image.new("RGB", (12, 8), "red").save(path)
        return path, payload

    monkeypatch.setattr(legacy, "_load_legacy_ref", replace_after_verification)

    summary = run_job(run)

    presentation = Presentation(summary["outputs"]["original"])
    picture = next(
        shape for shape in presentation.slides[0].shapes
        if shape.shape_type == 13
    )
    with Image.open(BytesIO(picture.image.blob)) as embedded:
        assert embedded.convert("RGB").getpixel((0, 0)) == (0, 128, 0)


def test_native_pdf_does_not_expose_mutable_image_snapshot_to_assembler(
    tmp_path: Path, monkeypatch,
) -> None:
    image = tmp_path / "asset.png"
    Image.new("RGB", (12, 8), "green").save(image)
    source = tmp_path / "native.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawImage(str(image), 20, 15, width=48, height=32)
    document.save()
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
    )
    original_assemble = image_to_ppt._assemble_prepared_slide

    def tamper_snapshot(slide_data, *args, **kwargs):
        component = slide_data["visual_elements"][0]["component"]
        if "path" in component:
            Image.new("RGB", (12, 8), "red").save(component["path"])
        return original_assemble(slide_data, *args, **kwargs)

    monkeypatch.setattr(image_to_ppt, "_assemble_prepared_slide", tamper_snapshot)

    summary = run_job(run)

    presentation = Presentation(summary["outputs"]["original"])
    picture = next(
        shape for shape in presentation.slides[0].shapes
        if shape.shape_type == 13
    )
    with Image.open(BytesIO(picture.image.blob)) as embedded:
        assert embedded.convert("RGB").getpixel((0, 0)) == (0, 128, 0)


def test_native_pdf_preserves_italic_text(tmp_path: Path) -> None:
    source = tmp_path / "italic.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.setFont("Helvetica-Oblique", 14)
    document.drawString(20, 70, "Italic")
    document.save()
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
    )

    summary = run_job(run)

    presentation = Presentation(summary["outputs"]["original"])
    textbox = next(
        shape for shape in presentation.slides[0].shapes
        if shape.has_text_frame and shape.text == "Italic"
    )
    assert textbox.text_frame.paragraphs[0].runs[0].font.italic is True


def test_native_pdf_writes_image_crop_and_character_spacing_xml(
    tmp_path: Path, monkeypatch,
) -> None:
    image = tmp_path / "asset.png"
    Image.new("RGB", (12, 8), "green").save(image)
    source = tmp_path / "native-crop-spacing.pdf"
    _write_cropped_spacing_pdf(source, image)

    def unexpected_worker(*args, **kwargs):
        raise AssertionError("native PDF must not create OCR or visual workers")

    monkeypatch.setattr(image_to_ppt, "create_ocr_worker_pool", unexpected_worker)
    monkeypatch.setattr(image_to_ppt, "create_visual_worker_pool", unexpected_worker)
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
    )

    summary = run_job(run)

    with ZipFile(summary["outputs"]["original"]) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml")
    assert b'<a:srcRect r="16667"/>' in slide_xml
    assert b'spc="900"' in slide_xml


def test_hybrid_pdf_bypasses_workers_and_keeps_editable_text(
    tmp_path: Path, monkeypatch,
) -> None:
    source = tmp_path / "hybrid.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawString(20, 95, "Editable heading")
    document.setFillColorRGB(0.1, 0.3, 0.8)
    document.roundRect(45, 30, 140, 30, 6, fill=1, stroke=0)
    document.setFillColorRGB(1, 1, 1)
    document.setFont("Times-Roman", 16)
    document.drawString(60, 35, "a / b = c")
    document.save()

    def unexpected_worker(*args, **kwargs):
        raise AssertionError("hybrid PDF must not create OCR or visual workers")

    monkeypatch.setattr(image_to_ppt, "create_ocr_worker_pool", unexpected_worker)
    monkeypatch.setattr(image_to_ppt, "create_visual_worker_pool", unexpected_worker)
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
        pipeline_mode="fast",
    )

    summary = run_job(run)

    presentation = Presentation(summary["outputs"]["original"])
    slide = presentation.slides[0]
    assert any(
        shape.has_text_frame and shape.text == "Editable heading"
        for shape in slide.shapes
    )
    assert any(
        shape.has_text_frame and shape.text == "a / b = c"
        for shape in slide.shapes
    )
    patches = [
        shape for shape in slide.shapes
        if shape.shape_type == 13 and shape.name.startswith("image2editable:patch-")
    ]
    assert len(patches) == 1
    assert patches[0].width * patches[0].height < (
        presentation.slide_width * presentation.slide_height * 0.35
    )
    state = RunStore.open(run).read_json(
        "pages/page_001/reconstruction/component_state.json"
    )
    assert state["route"] == "pdf_native"


def test_mirrored_transparent_pdf_image_bypasses_workers(
    tmp_path: Path, monkeypatch,
) -> None:
    image = tmp_path / "transparent.png"
    pixels = Image.new("RGBA", (12, 8), (255, 0, 0, 0))
    pixels.putpixel((0, 0), (0, 255, 0, 255))
    pixels.save(image)
    source = tmp_path / "mirrored.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.saveState()
    document.translate(80, 20)
    document.scale(-1, 1)
    document.drawImage(str(image), 0, 0, width=40, height=30, mask="auto")
    document.restoreState()
    document.save()

    def unexpected_worker(*args, **kwargs):
        raise AssertionError("mirrored PDF image must not create workers")

    monkeypatch.setattr(image_to_ppt, "create_ocr_worker_pool", unexpected_worker)
    monkeypatch.setattr(image_to_ppt, "create_visual_worker_pool", unexpected_worker)
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
        pipeline_mode="fast",
    )

    summary = run_job(run)

    presentation = Presentation(summary["outputs"]["original"])
    picture = next(
        shape for shape in presentation.slides[0].shapes
        if shape.shape_type == 13
    )
    with Image.open(BytesIO(picture.image.blob)) as embedded:
        assert embedded.mode == "RGBA"
        assert embedded.getpixel((embedded.width - 1, 0)) == (0, 255, 0, 255)


def test_near_axis_aligned_pdf_image_bypasses_workers(
    tmp_path: Path, monkeypatch,
) -> None:
    image = tmp_path / "asset.png"
    Image.new("RGB", (12, 8), "green").save(image)
    source = tmp_path / "near-axis.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.saveState()
    document.transform(1, 0.0000008, 0, 1, 0, 0)
    document.drawImage(str(image), 20, 15, width=40, height=30)
    document.restoreState()
    document.save()

    def unexpected_worker(*args, **kwargs):
        raise AssertionError("near-axis PDF image must not create workers")

    monkeypatch.setattr(image_to_ppt, "create_ocr_worker_pool", unexpected_worker)
    monkeypatch.setattr(image_to_ppt, "create_visual_worker_pool", unexpected_worker)
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
        pipeline_mode="fast",
    )

    summary = run_job(run)

    presentation = Presentation(summary["outputs"]["original"])
    assert sum(
        shape.shape_type == 13 for shape in presentation.slides[0].shapes
    ) == 1


def test_native_pdf_writes_fill_opacity_xml(tmp_path: Path) -> None:
    source = tmp_path / "opacity.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.setFillAlpha(0.5)
    document.setFillColorRGB(1, 0, 0)
    document.rect(20, 20, 40, 30, fill=1, stroke=0)
    document.setFillColorRGB(0, 0, 0)
    document.drawString(80, 70, "Transparent")
    document.setStrokeAlpha(0.25)
    document.setStrokeColorRGB(0, 0, 1)
    document.line(20, 90, 60, 90)
    document.save()
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=tmp_path / "output.pptx",
        slide_size="original",
        pipeline_mode="fast",
    )

    summary = run_job(run)

    with ZipFile(summary["outputs"]["original"]) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml")
    assert slide_xml.count(b'<a:alpha val="50000"/>') == 2
    assert slide_xml.count(b'<a:alpha val="25000"/>') == 1


def test_native_pdf_failure_after_assembly_keeps_retry_assets(
    tmp_path: Path, monkeypatch,
) -> None:
    image = tmp_path / "asset.png"
    Image.new("RGB", (12, 8), "green").save(image)
    source = tmp_path / "native.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawImage(str(image), 20, 15, width=48, height=32)
    document.drawString(80, 70, "Editable")
    document.save()
    output = tmp_path / "output.pptx"
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=output,
        slide_size="original",
    )

    original_summary = runtime._performance_summary
    monkeypatch.setattr(
        runtime,
        "_performance_summary",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("summary failed")),
    )

    with pytest.raises(RuntimeError, match="summary failed"):
        runtime.run_job(run)

    assert (run / "pages/page_001/source.png").is_file()
    assert (run / "pages/page_001/pdf-assets/image-0001.png").is_file()
    assert not output.exists()

    monkeypatch.setattr(runtime, "_performance_summary", original_summary)
    assert runtime.retry_page(run, "page_001")["run"]["status"] == "prepared"
    assert runtime.run_job(run)["status"] == "completed"


def test_native_pdf_compensation_preserves_concurrently_replaced_output(
    tmp_path: Path, monkeypatch,
) -> None:
    source = tmp_path / "native.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawString(20, 70, "Editable")
    document.save()
    output = tmp_path / "output.pptx"
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=output,
        slide_size="original",
    )

    def replace_output_then_fail(*args, **kwargs):
        output.unlink()
        output.write_bytes(b"concurrent replacement")
        raise RuntimeError("summary failed")

    monkeypatch.setattr(runtime, "_performance_summary", replace_output_then_fail)

    with pytest.raises(RuntimeError, match="summary failed") as error:
        runtime.run_job(run)

    assert error.value.__cause__ is not None
    assert output.read_bytes() == b"concurrent replacement"
    summary = RunStore.open(run).read_json("run_summary.json")
    assert summary["retry_blocked"] is True
    with pytest.raises(RuntimeError, match="blocked"):
        runtime.retry_page(run, "page_001")


def test_native_pdf_delivery_replacement_blocks_retry(
    tmp_path: Path, monkeypatch,
) -> None:
    source = tmp_path / "native.pdf"
    document = canvas.Canvas(str(source), pagesize=(200, 120))
    document.drawString(20, 70, "Editable")
    document.save()
    output = tmp_path / "output.pptx"
    run = prepare_pdf_job(
        source,
        run_dir=tmp_path / "run",
        output_path=output,
        slide_size="original",
    )
    original_record = legacy._record_legacy_delivery

    def replace_before_delivery(*args, **kwargs):
        output.unlink()
        output.write_bytes(b"concurrent replacement")
        return original_record(*args, **kwargs)

    monkeypatch.setattr(
        legacy, "_record_legacy_delivery", replace_before_delivery
    )

    with pytest.raises(RuntimeError, match="Legacy output identity changed"):
        runtime.run_job(run)

    assert output.read_bytes() == b"concurrent replacement"
    summary = RunStore.open(run).read_json("run_summary.json")
    assert summary["retry_blocked"] is True
    with pytest.raises(RuntimeError, match="blocked"):
        runtime.retry_page(run, "page_001")
