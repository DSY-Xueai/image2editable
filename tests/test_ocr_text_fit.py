from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches


@pytest.fixture(params=["scripts", "skills/image-to-ppt/scripts"])
def assembler(request):
    path = Path(__file__).resolve().parents[1] / request.param / "ppt_assemble.py"
    spec = importlib.util.spec_from_file_location("ocr_fit_assembler", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_ocr_text_fits_width_without_moving_box(assembler, monkeypatch, rotation):
    font_path = Path(__file__).resolve().parents[1] / "benchmarks/release/fonts/NotoSansSC[wght].ttf"
    measured_font = ImageFont.truetype(str(font_path), 1000)
    monkeypatch.setattr(ImageFont, "truetype", lambda *a, **kw: measured_font)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = [100, 100, 60, 20] if rotation in {0, 180} else [120, 80, 20, 60]
    text = "\u5ba2\u6237\u51c0\u63a8\u8350\u503c"
    shape = assembler._add_textbox(
        slide, {"box": box, "text": text, "font_size": 30, "rotation": rotation},
        960, 540, assembler.compute_slide_transform(960, 540, "original"),
    )
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.text == text
    assert measured_font.getlength(text) * run.font.size.pt / 1000 <= shape.width.pt + 0.01
    assert abs(shape.width - Inches(60 / 72)) <= 1
    assert (shape.left + shape.width / 2) / 12700 == pytest.approx(130, abs=0.01)
    assert (shape.top + shape.height / 2) / 12700 == pytest.approx(110, abs=0.01)


@pytest.mark.parametrize("native", [False, True])
def test_fitting_keeps_small_ocr_and_native_pdf_font_size(assembler, monkeypatch, native):
    measured_font = ImageFont.load_default(size=1000)
    monkeypatch.setattr(ImageFont, "truetype", lambda *a, **kw: measured_font)
    deck = Presentation()
    item = {"box": [10, 10, 100, 30], "text": "Label", "font_size": 8}
    if native:
        item.update(font_size_pt=60, character_spacing_pt=2)
    shape = assembler._add_textbox(
        deck.slides.add_slide(deck.slide_layouts[6]), item,
        960, 540, assembler.compute_slide_transform(960, 540, "original"),
    )
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(60 if native else 8, abs=0.011)
    if native:
        assert run._r.get_or_add_rPr().get("spc") == "200"


def test_missing_measurement_font_does_not_abort_conversion(assembler, monkeypatch):
    def unavailable(*args, **kwargs):
        raise OSError("font is not installed")

    monkeypatch.setattr(ImageFont, "truetype", unavailable)
    assert assembler._fit_ocr_font_size("Test", "Arial", False, False, 18, 1) == 18


def test_narrow_ocr_frame_uses_valid_font_size_and_wraps_in_place(assembler, monkeypatch, tmp_path):
    measured_font = ImageFont.load_default(size=1000)
    monkeypatch.setattr(ImageFont, "truetype", lambda *a, **kw: measured_font)
    deck = Presentation()
    text = "This is a long narrow OCR label"
    shape = assembler._add_textbox(
        deck.slides.add_slide(deck.slide_layouts[6]),
        {"box": [10, 10, 10, 20], "text": text, "font": "Arial", "font_size": 12},
        960, 540, assembler.compute_slide_transform(960, 540, "original"),
    )
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.text == text
    assert run.font.size.pt == 1
    assert shape.text_frame.word_wrap is True
    assert shape.text_frame._txBody.bodyPr.normAutofit is not None
    path = tmp_path / "narrow.pptx"
    deck.save(path)
    reopened = Presentation(path).slides[0].shapes[0]
    assert reopened.text == text
    assert reopened.width == shape.width
    assert reopened.height == shape.height


def test_measurement_font_is_loaded_once_for_repeated_labels(assembler, monkeypatch):
    calls = []
    font = ImageFont.load_default(size=1000)

    def load(*args, **kwargs):
        calls.append(args)
        return font

    monkeypatch.setattr(ImageFont, "truetype", load)
    for text in ("Short", "Longer text", "Another label"):
        assembler._fit_ocr_font_size(text, "Arial", True, False, 18, 1)
    assert len(calls) == 1
