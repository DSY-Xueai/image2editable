from pathlib import Path
import shutil

from PIL import Image
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches
import pytest

from image2editable.native_pdf_quality import evaluate_native_pdf_page


@pytest.fixture
def native_candidate(tmp_path: Path):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    pptx = tmp_path / "output.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(.5))
    shape.name = "image2editable:text-0001"
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Editable text"
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0, 0, 0)
    presentation.save(pptx)
    source, rendered = tmp_path / "source.png", tmp_path / "rendered.png"
    # Unit-level renderer inputs; actual Office output is tested separately.
    Image.new("RGB", (720, 360), "white").save(source)
    Image.new("RGB", (720, 360), "white").save(rendered)
    return dict(
        pptx_path=pptx, page_number=1, source_path=source, rendered_path=rendered,
        analysis={"width_pt": 720, "height_pt": 360, "objects": [{
            "id": "text-0001", "z_index": 1, "type": "text", "text": "Editable text",
            "font": "Arial", "font_size": 20, "color_rgb": [0, 0, 0],
            "bbox_pt": [72, 252, 288, 288],
        }]},
    )


def test_native_quality_accepts_matching_objects_and_pixels(native_candidate):
    result = evaluate_native_pdf_page(**native_candidate)
    assert result["report"]["accepted"]
    assert result["report"]["violations"] == []
    assert len(result["pptx_sha256"]) == 64


@pytest.mark.parametrize("change", ["missing", "duplicate", "text", "font", "transparent", "position"])
def test_native_quality_rejects_corrupted_editable_objects(native_candidate, change):
    path = native_candidate["pptx_path"]
    presentation = Presentation(path)
    slide = presentation.slides[0]
    shape = slide.shapes[0]
    if change == "missing":
        shape._element.getparent().remove(shape._element)
    elif change == "duplicate":
        from copy import deepcopy
        slide.shapes._spTree.append(deepcopy(shape._element))
    elif change == "text":
        shape.text = "Wrong text"
    elif change == "font":
        shape.text_frame.paragraphs[0].runs[0].font.name = "Times New Roman"
    elif change == "transparent":
        color = shape.text_frame.paragraphs[0].runs[0]._r.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        alpha = OxmlElement("a:alpha")
        alpha.set("val", "0")
        color.append(alpha)
    else:
        shape.left += Inches(.1)
    presentation.save(path)
    assert not evaluate_native_pdf_page(**native_candidate)["report"]["accepted"]


def test_native_quality_uses_existing_visual_thresholds(native_candidate):
    Image.new("RGB", (720, 360), "black").save(native_candidate["rendered_path"])
    result = evaluate_native_pdf_page(**native_candidate)
    assert not result["report"]["accepted"]
    assert "visual_difference" in result["report"]["violations"]


@pytest.mark.parametrize("tamper", [None, "pptx", "render", "source", "report", "missing"])
def test_release_native_validation_recomputes_bound_evidence(native_candidate, tmp_path, tamper):
    from image2editable.inputs import sha256_file
    from image2editable.store import RunStore
    from scripts import release_benchmark as benchmark

    store = RunStore(tmp_path / "run")
    relative = "pages/page_001/reconstruction/"
    directory = store.root / relative
    store.write_json(relative + "native_page.json", {
        "page_id": "page_001", "analysis": native_candidate["analysis"],
    })
    analysis_ref = {"path": relative + "native_page.json", "sha256": sha256_file(directory / "native_page.json")}
    store.write_json(relative + "component_state.json", {
        "route": "pdf_native", "native_page_ref": analysis_ref,
    })
    shutil.copyfile(native_candidate["source_path"], directory / "native-source.png")
    shutil.copyfile(native_candidate["rendered_path"], directory / "native-render-original.png")
    quality = evaluate_native_pdf_page(**native_candidate)
    quality.update(page_id="page_001", analysis_sha256=analysis_ref["sha256"], renderer="libreoffice")
    if tamper == "report":
        quality["report"]["visual_metrics"]["mae"] = 1.0
    store.write_json(relative + "native-quality.json", {"page_id": "page_001", "variants": {"original": quality}})
    store.write_json(relative + "component_delivery.json", {
        "page_id": "page_001",
        "native_quality_ref": {"path": relative + "native-quality.json", "sha256": sha256_file(directory / "native-quality.json")},
        "outputs": {"original": {"path": str(native_candidate["pptx_path"]), "sha256": quality["pptx_sha256"]}},
    })
    if tamper == "pptx":
        presentation = Presentation(native_candidate["pptx_path"])
        presentation.slides[0].shapes[0].text = "Changed after QA"
        presentation.save(native_candidate["pptx_path"])
    elif tamper in {"source", "render"}:
        name = "native-source.png" if tamper == "source" else "native-render-original.png"
        Image.new("RGB", (720, 360), "black").save(directory / name)
    elif tamper == "missing":
        (directory / "native-quality.json").unlink()
    case = {"kind": "pdf", "expected_pages": [{
        "expected_status": "validated", "min_text_boxes": 1, "min_visual_components": 0,
    }]}
    result = benchmark.BenchmarkCaseResult("pdf-one", str(store.root), [{"page_id": "page_001", "status": "validated"}], 0)
    if tamper is None:
        benchmark._validate_batch_case(case, result)
    else:
        with pytest.raises(benchmark.BenchmarkFailure, match="invalid_quality_result"):
            benchmark._validate_batch_case(case, result)
